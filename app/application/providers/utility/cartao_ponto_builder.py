"""Monta o arquivo de cartoes de ponto no formato usado para impressao e assinatura.

Uma aba por funcionario, uma linha por dia do periodo — inclusive dias sem
batida e finais de semana, porque o cartao e assinado no papel. Os horarios
saem como texto HH:MM para imprimir exatamente o que foi calculado, sem
depender de formatacao de celula ou do locale de quem abrir o arquivo.
"""

from __future__ import annotations

from dataclasses import dataclass
from datetime import date, time as Time, timedelta
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side

from app.domain.services.rh_cartao_ponto import LinhaCartao

_COLUNAS = ["Data", "Dia", "entrada", "saída intervalo", "retorno intervalo", "saída", "assinatura"]
_LARGURAS = [12, 8, 12, 18, 20, 12, 30]
_DIAS_SEMANA = ["SEG", "TER", "QUA", "QUI", "SEX", "SAB", "DOM"]
_MESES = [
    "Janeiro", "Fevereiro", "Março", "Abril", "Maio", "Junho",
    "Julho", "Agosto", "Setembro", "Outubro", "Novembro", "Dezembro",
]
_LINHA_CABECALHO_TABELA = 6
_CARACTERES_PROIBIDOS_ABA = set('[]:*?/\\')
_LIMITE_NOME_ABA = 31


@dataclass(frozen=True)
class DiaCartao:
    data: date
    linha: LinhaCartao


@dataclass(frozen=True)
class CartaoFuncionario:
    nome: str
    cargo: str
    dias: list[DiaCartao]


class CartaoPontoBuilder:
    def build(
        self,
        empresa: str,
        inicio: date,
        fim: date,
        cartoes: list[CartaoFuncionario],
    ) -> bytes:
        workbook = Workbook()
        workbook.remove(workbook.active)
        usados: set[str] = set()

        for cartao in cartoes:
            sheet = workbook.create_sheet(self._nome_aba(cartao.nome, usados))
            self._build_cabecalho(sheet, empresa, cartao, inicio, fim)
            self._build_tabela(sheet, cartao)

        if not workbook.sheetnames:
            workbook.create_sheet("Sem funcionarios")

        buffer = BytesIO()
        workbook.save(buffer)
        return buffer.getvalue()

    def _nome_aba(self, nome: str, usados: set[str]) -> str:
        # O Excel limita nomes de aba a 31 caracteres e proibe alguns simbolos.
        limpo = "".join(ch for ch in nome if ch not in _CARACTERES_PROIBIDOS_ABA).strip() or "Funcionario"
        base = limpo[:_LIMITE_NOME_ABA]
        candidato = base
        sufixo = 2
        while candidato in usados:
            corte = _LIMITE_NOME_ABA - len(f" {sufixo}")
            candidato = f"{base[:corte]} {sufixo}"
            sufixo += 1
        usados.add(candidato)
        return candidato

    def _build_cabecalho(
        self,
        sheet,
        empresa: str,
        cartao: CartaoFuncionario,
        inicio: date,
        fim: date,
    ) -> None:
        sheet["A1"] = empresa
        sheet["A1"].font = Font(bold=True)
        sheet["A2"] = f"Nome-{cartao.nome}"
        sheet["A3"] = f"Função-{cartao.cargo}"
        sheet["A4"] = self._referencia(inicio, fim)

    def _referencia(self, inicio: date, fim: date) -> str:
        mes_inteiro = (
            inicio.day == 1
            and inicio.month == fim.month
            and inicio.year == fim.year
            and (fim + timedelta(days=1)).month != fim.month
        )
        if mes_inteiro:
            return f"Ref: {_MESES[inicio.month - 1]} {inicio.year}"
        return f"Período: {inicio.strftime('%d/%m/%Y')} a {fim.strftime('%d/%m/%Y')}"

    def _build_tabela(self, sheet, cartao: CartaoFuncionario) -> None:
        fina = Side(style="thin")
        borda = Border(left=fina, right=fina, top=fina, bottom=fina)
        centro = Alignment(horizontal="center")

        for indice, titulo in enumerate(_COLUNAS, start=1):
            celula = sheet.cell(row=_LINHA_CABECALHO_TABELA, column=indice, value=titulo)
            celula.font = Font(bold=True)
            celula.border = borda
            celula.alignment = centro
            sheet.column_dimensions[celula.column_letter].width = _LARGURAS[indice - 1]

        for offset, dia in enumerate(cartao.dias):
            linha = _LINHA_CABECALHO_TABELA + 1 + offset
            valores = [
                dia.data.strftime("%d/%m/%Y"),
                _DIAS_SEMANA[dia.data.weekday()],
                self._hhmm(dia.linha.entrada),
                self._hhmm(dia.linha.saida_intervalo),
                self._hhmm(dia.linha.retorno_intervalo),
                self._hhmm(dia.linha.saida),
                None,
            ]
            for indice, valor in enumerate(valores, start=1):
                celula = sheet.cell(row=linha, column=indice, value=valor)
                celula.border = borda
                if indice <= 6:
                    celula.alignment = centro

    def _hhmm(self, valor: Time | None) -> str | None:
        return valor.strftime("%H:%M") if valor is not None else None
