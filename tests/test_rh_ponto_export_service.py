from datetime import date, datetime, timezone
from io import BytesIO
from types import SimpleNamespace
from uuid import uuid4

import pytest
from openpyxl import load_workbook

from app.application.services.rh_ponto_export_service import RhPontoExportService
from app.domain.entities.rh import RegistroPonto, StatusPonto, TipoPonto
from app.domain.entities.user import Roles
from app.domain.errors import DomainError


class _FakeFuncionarioRepo:
    def __init__(self, funcionarios):
        self._funcionarios = funcionarios

    async def get_by_id(self, id, team_id):
        return next(f for f in self._funcionarios if f.id == id)

    async def list_active_by_team(self, team_id, limit, offset):
        return self._funcionarios if offset == 0 else []


class _FakeRegistroRepo:
    def __init__(self, registros):
        self._registros = registros
        self.ultima_janela = None

    async def list_by_competencia(self, team_id, funcionario_ids, start, end):
        self.ultima_janela = (start, end)
        return [r for r in self._registros if r.funcionario_id in funcionario_ids]


def _funcionario(nome, cargo):
    return SimpleNamespace(id=uuid4(), nome=nome, cargo=cargo)


def _user(team_id, role=Roles.ADMIN):
    return SimpleNamespace(
        id=uuid4(),
        role=role,
        team=SimpleNamespace(id=team_id, title="Arcaika Ltda", cnpj="00000000000100"),
    )


def _reg(funcionario_id, team_id, momento, tipo):
    return RegistroPonto(
        team_id=team_id,
        funcionario_id=funcionario_id,
        tipo=tipo,
        timestamp=momento,
        latitude=0.0,
        longitude=0.0,
        status=StatusPonto.VALIDADO,
    )


def _service(funcionarios, registros):
    return RhPontoExportService(
        funcionario_repo=_FakeFuncionarioRepo(funcionarios),
        registro_ponto_repo=_FakeRegistroRepo(registros),
    )


@pytest.mark.asyncio
async def test_exporta_uma_aba_por_funcionario_ativo():
    team_id = uuid4()
    funcionarios = [_funcionario("Sandro Barbosa", "Serralheiro"), _funcionario("Ana Lima", "Pedreira")]
    service = _service(funcionarios, [])
    conteudo = await service.exportar_cartoes(
        _user(team_id), date(2026, 3, 1), date(2026, 3, 31), funcionario_id=None
    )
    nomes = load_workbook(BytesIO(conteudo)).sheetnames
    assert nomes == ["Sandro Barbosa", "Ana Lima"]


@pytest.mark.asyncio
async def test_filtra_por_funcionario_quando_informado():
    team_id = uuid4()
    funcionarios = [_funcionario("Sandro Barbosa", "Serralheiro"), _funcionario("Ana Lima", "Pedreira")]
    service = _service(funcionarios, [])
    conteudo = await service.exportar_cartoes(
        _user(team_id), date(2026, 3, 1), date(2026, 3, 31), funcionario_id=funcionarios[1].id
    )
    assert load_workbook(BytesIO(conteudo)).sheetnames == ["Ana Lima"]


@pytest.mark.asyncio
async def test_batida_noturna_cai_no_dia_local_correto():
    team_id = uuid4()
    funcionario = _funcionario("Sandro Barbosa", "Serralheiro")
    # 10/03 as 18:00 e 23:00 locais; a saida em UTC ja e 11/03 as 02:00.
    registros = [
        _reg(funcionario.id, team_id, datetime(2026, 3, 10, 21, 0, tzinfo=timezone.utc), TipoPonto.ENTRADA),
        _reg(funcionario.id, team_id, datetime(2026, 3, 11, 2, 0, tzinfo=timezone.utc), TipoPonto.SAIDA),
    ]
    service = _service([funcionario], registros)
    conteudo = await service.exportar_cartoes(
        _user(team_id), date(2026, 3, 1), date(2026, 3, 31), funcionario_id=funcionario.id
    )
    aba = load_workbook(BytesIO(conteudo)).worksheets[0]
    # Linha 6 e o cabecalho; 01/03 e a linha 7, entao 10/03 e a linha 16.
    assert aba["A16"].value == "10/03/2026"
    assert aba["C16"].value == "18:00"
    assert aba["F16"].value == "23:00"


@pytest.mark.asyncio
async def test_recusa_periodo_invertido():
    service = _service([], [])
    with pytest.raises(DomainError):
        await service.exportar_cartoes(
            _user(uuid4()), date(2026, 3, 31), date(2026, 3, 1), funcionario_id=None
        )


@pytest.mark.asyncio
async def test_recusa_periodo_maior_que_um_ano():
    service = _service([], [])
    with pytest.raises(DomainError):
        await service.exportar_cartoes(
            _user(uuid4()), date(2025, 1, 1), date(2026, 6, 1), funcionario_id=None
        )


@pytest.mark.asyncio
async def test_recusa_usuario_sem_perfil_de_rh():
    service = _service([], [])
    with pytest.raises(DomainError):
        await service.exportar_cartoes(
            _user(uuid4(), role=Roles.FUNCIONARIO), date(2026, 3, 1), date(2026, 3, 31), funcionario_id=None
        )


def test_nome_arquivo_usa_o_periodo():
    service = _service([], [])
    assert service.nome_arquivo(date(2026, 3, 1), date(2026, 3, 31)) == "cartoes-ponto-2026-03-01-a-2026-03-31.xlsx"
