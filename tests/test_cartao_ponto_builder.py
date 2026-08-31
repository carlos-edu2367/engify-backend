from datetime import date, time
from decimal import Decimal
from io import BytesIO

from openpyxl import load_workbook

from app.application.providers.utility.cartao_ponto_builder import (
    CartaoFuncionario,
    CartaoPontoBuilder,
    DiaCartao,
)
from app.domain.services.rh_cartao_ponto import LinhaCartao

CABECALHO = ["Data", "Dia", "entrada", "saída intervalo", "retorno intervalo", "saída", "HE", "atrasos"]


def _cartao_marco():
    dias = [
        DiaCartao(
            data=date(2026, 3, 1),
            linha=LinhaCartao(None, None, None, None),
        ),
        DiaCartao(
            data=date(2026, 3, 2),
            linha=LinhaCartao(time(8, 0), time(12, 0), time(13, 0), time(17, 48)),
            extra_min=Decimal("48"),
        ),
    ]
    return CartaoFuncionario(nome="Sandro Barbosa", cargo="Serralheiro", dias=dias)


def _abrir(conteudo: bytes):
    return load_workbook(BytesIO(conteudo))


def test_cabecalho_traz_empresa_nome_funcao_e_referencia():
    conteudo = CartaoPontoBuilder().build(
        empresa="Arcaika Ltda — CNPJ 00.000.000/0001-00",
        inicio=date(2026, 3, 1),
        fim=date(2026, 3, 31),
        cartoes=[_cartao_marco()],
    )
    aba = _abrir(conteudo).worksheets[0]
    assert aba["A1"].value == "Arcaika Ltda — CNPJ 00.000.000/0001-00"
    assert aba["A2"].value == "Nome-Sandro Barbosa"
    assert aba["A3"].value == "Função-Serralheiro"
    assert aba["A4"].value == "Ref: Março 2026"


def test_periodo_parcial_usa_rotulo_de_periodo():
    conteudo = CartaoPontoBuilder().build(
        empresa="Arcaika Ltda",
        inicio=date(2026, 3, 5),
        fim=date(2026, 3, 20),
        cartoes=[_cartao_marco()],
    )
    aba = _abrir(conteudo).worksheets[0]
    assert aba["A4"].value == "Período: 05/03/2026 a 20/03/2026"


def test_tabela_tem_o_cabecalho_do_modelo_na_linha_6():
    conteudo = CartaoPontoBuilder().build(
        empresa="Arcaika Ltda", inicio=date(2026, 3, 1), fim=date(2026, 3, 31), cartoes=[_cartao_marco()]
    )
    aba = _abrir(conteudo).worksheets[0]
    assert [aba.cell(row=6, column=c).value for c in range(1, 9)] == CABECALHO


def test_dia_com_batidas_sai_como_texto_hh_mm():
    conteudo = CartaoPontoBuilder().build(
        empresa="Arcaika Ltda", inicio=date(2026, 3, 1), fim=date(2026, 3, 31), cartoes=[_cartao_marco()]
    )
    aba = _abrir(conteudo).worksheets[0]
    assert aba["A8"].value == "02/03/2026"
    assert aba["B8"].value == "SEG"
    assert [aba.cell(row=8, column=c).value for c in range(3, 7)] == ["08:00", "12:00", "13:00", "17:48"]
    # HE do dia: 48 minutos de hora extra
    assert aba["G8"].value == "00:48"
    assert aba["H8"].value == "00:00"


def test_dia_sem_batidas_deixa_celulas_de_horario_vazias():
    conteudo = CartaoPontoBuilder().build(
        empresa="Arcaika Ltda", inicio=date(2026, 3, 1), fim=date(2026, 3, 31), cartoes=[_cartao_marco()]
    )
    aba = _abrir(conteudo).worksheets[0]
    assert aba["A7"].value == "01/03/2026"
    assert aba["B7"].value == "DOM"
    assert [aba.cell(row=7, column=c).value for c in range(3, 7)] == [None, None, None, None]
    # sem turno e sem batida: nao ha hora extra nem atraso a marcar
    assert aba["G7"].value == "00:00"
    assert aba["H7"].value == "00:00"


def test_coluna_de_atrasos_formata_minutos_como_hh_mm():
    dias = [
        DiaCartao(
            data=date(2026, 3, 3),
            linha=LinhaCartao(time(8, 0), None, None, time(15, 30)),
            falta_min=Decimal("90"),
        ),
    ]
    cartao = CartaoFuncionario(nome="Sandro Barbosa", cargo="Serralheiro", dias=dias)
    conteudo = CartaoPontoBuilder().build(
        empresa="Arcaika Ltda", inicio=date(2026, 3, 1), fim=date(2026, 3, 31), cartoes=[cartao]
    )
    aba = _abrir(conteudo).worksheets[0]
    assert aba["G7"].value == "00:00"
    assert aba["H7"].value == "01:30"


def test_uma_aba_por_funcionario_com_nome_truncado_e_sem_colisao():
    longo = "Maria Aparecida da Conceicao Nascimento Silva"
    cartoes = [
        CartaoFuncionario(nome=longo, cargo="Pedreiro", dias=[]),
        CartaoFuncionario(nome=longo, cargo="Servente", dias=[]),
    ]
    conteudo = CartaoPontoBuilder().build(
        empresa="Arcaika Ltda", inicio=date(2026, 3, 1), fim=date(2026, 3, 31), cartoes=cartoes
    )
    nomes = _abrir(conteudo).sheetnames
    assert len(nomes) == 2
    assert nomes[0] != nomes[1]
    assert all(len(nome) <= 31 for nome in nomes)
